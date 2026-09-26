"""Counts, proportions, statistics, Excel workbook and plots."""

import datetime
import json
import os
import platform

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402

from . import __version__, stats  # noqa: E402

__all__ = ["build_tables", "write_outputs", "experiment_name"]

# Filename fields not used to identify an organoid (they may differ between slices)
NON_ORGANOID_FIELDS = ("slice", "objective", "imaging")


def _title(field):
    return field[:1].upper() + field[1:]


def experiment_name(per_image, fields):
    """Diff_Day_Immuno style name from fields that are constant across the run."""
    parts = []
    for f in ("diff", "day", "immuno"):
        col = _title(f)
        if f in fields and col in per_image and per_image[col].nunique() == 1:
            parts.append(str(per_image[col].iloc[0]))
    return "_".join(parts) or "NucleiQuant"


def build_tables(project, summaries):
    """Per-image / per-organoid tables from per-image summaries.

    summaries: list of dicts {"image", "counts": {category_id: n}}.
    Returns a dict of DataFrames and info.
    """
    cats = project.categories
    names = {c["id"]: c["name"] for c in cats}
    fields = list(project.pattern().groupindex)
    rows = []
    for s in summaries:
        try:
            meta = project.parse(s["image"])
        except ValueError:
            meta = {}
        row = {"Image_name": s["image"]}
        for f in fields:
            row[_title(f)] = meta.get(f, "")
            if f == "clone":
                row["Genotype"] = project.data["genotypes"].get(meta.get("clone", ""), "")
        total = 0
        for c in cats:
            n = int(s["counts"].get(c["id"], 0))
            row[c["name"]] = n
            total += n
        row["Total"] = total
        row["Living"] = total - int(s["counts"].get("dead", 0))
        rows.append(row)
    per_image = pd.DataFrame(rows)
    cat_cols = [names[c["id"]] for c in cats]
    info = {"fields": fields, "category_columns": cat_cols, "warnings": []}

    if "organoid" in fields and len(per_image):
        key = [_title(f) for f in fields if f not in NON_ORGANOID_FIELDS]
        if "clone" in fields:
            key.insert(key.index("Clone") + 1, "Genotype")
        slice_col = "Slice" if "slice" in fields else None
        agg = {c: "sum" for c in cat_cols + ["Total", "Living"]}
        per_org = per_image.groupby(key, dropna=False, sort=True).agg(agg).reset_index()
        n_sl = per_image.groupby(key, dropna=False, sort=True)[slice_col or "Image_name"].nunique().reset_index(name="Slices")
        per_org = per_org.merge(n_sl, on=key)
        info["unit"] = "organoid"
    else:
        key = ["Image_name"] + [_title(f) for f in fields]
        per_org = per_image.copy()
        per_org["Slices"] = 1
        info["unit"] = "image"
        info["warnings"].append("No 'organoid' field in the file names: each image is treated as one sample.")
    info["key"] = key
    min_slices = int(project.settings["min_slices"])
    curated = per_org[per_org["Slices"] >= min_slices].copy() if info["unit"] == "organoid" else per_org.copy()
    info["excluded"] = per_org[per_org["Slices"] < min_slices] if info["unit"] == "organoid" else per_org.iloc[0:0]

    def proportions(tab):
        out = tab[key + ["Slices", "Total", "Living"]].copy()
        for c in cats:
            n = names[c["id"]]
            with np.errstate(divide="ignore", invalid="ignore"):
                out[f"{n} (% all)"] = 100 * tab[n] / tab["Total"].where(tab["Total"] > 0)
                if c["id"] != "dead":
                    out[f"{n} (% living)"] = 100 * tab[n] / tab["Living"].where(tab["Living"] > 0)
        return out

    props = proportions(curated)
    props_all = proportions(per_org)
    return {"per_image": per_image, "per_organoid": per_org, "curated": curated,
            "proportions": props, "proportions_all_organoids": props_all, "info": info}


def run_stats(project, tables):
    """Group comparisons on the curated per-organoid proportions."""
    props = tables["proportions"]
    group_col = "Genotype" if project.settings.get("group_by", "genotype") == "genotype" else "Clone"
    if group_col not in props or (props[group_col] == "").all():
        group_col = "Clone" if "Clone" in props else None
    if group_col is None:
        return pd.DataFrame(), ["No clone/genotype field in the file names: no groups to compare."], None
    cats = project.categories
    living = [f"{c['name']} (% living)" for c in cats if c["id"] != "dead"]
    allc = [f"{c['name']} (% all)" for c in cats]
    rows1, warn1 = stats.compare_groups(props, group_col, living, "% of living cells")
    rows2, _ = stats.compare_groups(props, group_col, allc, "% of all nuclei")
    df = pd.DataFrame(rows1 + rows2)
    if len(df):
        df["Category"] = df["Category"].str.replace(r" \(% (living|all)\)$", "", regex=True)
    return df, warn1, group_col


def _autosize(ws):
    from openpyxl.styles import Font, PatternFill
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EDEEF0")
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
    ws.freeze_panes = "A2"


def write_outputs(project, summaries, classifier_info=None):
    """Write the Excel workbook and plots. Returns a JSON-able summary for the app."""
    tables = build_tables(project, summaries)
    stats_df, stat_warnings, group_col = run_stats(project, tables)
    info = tables["info"]
    fields = info["fields"]
    name = experiment_name(tables["per_image"], fields)
    out_dir = os.path.join(project.dir, "results")
    os.makedirs(out_dir, exist_ok=True)
    xlsx = os.path.join(out_dir, f"{name}_counts.xlsx")

    cats = project.categories
    classifier_rows = []
    if classifier_info:
        cv = classifier_info.get("cross_validation", {})
        classifier_rows += [
            ("Trained", classifier_info.get("trained")),
            ("Labeled nuclei", classifier_info.get("n_labels")),
            ("Features per nucleus", classifier_info.get("n_features")),
            ("Trees", classifier_info.get("n_trees")),
            ("Validation", cv.get("scheme")),
            ("Accuracy on unseen crops", cv.get("accuracy")),
            ("Balanced accuracy", cv.get("balanced_accuracy")),
        ]
        for c in cats:
            n = classifier_info.get("labels_per_category", {}).get(c["id"], 0)
            rec = cv.get("recall", {}).get(c["id"])
            classifier_rows.append((f"{c['name']}: labels / recall", f"{n} / {rec:.2f}" if rec is not None else f"{n} / -"))
        for tf in classifier_info.get("top_features", [])[:15]:
            classifier_rows.append((f"Important feature: {tf['feature']}", round(tf["importance"], 4)))
    s = project.settings
    settings_rows = [
        ("NucleiQuant version", __version__),
        ("Date", datetime.datetime.now().astimezone().isoformat(timespec="seconds")),
        ("Images folder", project.images_dir),
        ("File name pattern", project.data["filename_template"]),
        ("Channels", ", ".join(f"{i + 1}: {c['name']}" for i, c in enumerate(project.data["channels"]))),
        ("Nuclear channel", project.data["channels"][project.data["nuclear_channel"]]["name"]),
        ("Segmentation", f"stardist_haug2, normalisation percentiles {s['norm_low']}-{s['norm_high']}"),
        ("Neighbourhood ring (px)", s["ring_radius"]),
        ("Perinuclear ring (px)", s["perinuclear_radius"]),
        ("Minimum slices per organoid", s["min_slices"]),
        ("Genotypes", json.dumps(project.data["genotypes"])),
        ("Python", platform.python_version()),
    ]

    with pd.ExcelWriter(xlsx, engine="openpyxl") as xw:
        tables["per_image"].to_excel(xw, sheet_name="per_image", index=False)
        tables["per_organoid"].to_excel(xw, sheet_name="per_organoid", index=False)
        tables["curated"].to_excel(xw, sheet_name="per_organoid_curated", index=False)
        tables["proportions"].round(3).to_excel(xw, sheet_name="proportions", index=False)
        if len(stats_df):
            stats_df.to_excel(xw, sheet_name="statistics", index=False)
        else:
            pd.DataFrame({"Note": stat_warnings or ["No statistics"]}).to_excel(xw, sheet_name="statistics", index=False)
        pd.DataFrame(classifier_rows, columns=["Item", "Value"]).to_excel(xw, sheet_name="classifier", index=False)
        pd.DataFrame(settings_rows, columns=["Setting", "Value"]).to_excel(xw, sheet_name="settings", index=False)
        for ws in xw.book.worksheets:
            _autosize(ws)

    plots = _plots(project, tables, group_col, out_dir)

    per_image = tables["per_image"]
    total = int(per_image["Total"].sum()) if len(per_image) else 0
    living = int(per_image["Living"].sum()) if len(per_image) else 0
    props = tables["proportions_all_organoids"]
    excluded_keys = set(map(tuple, info["excluded"][info["key"]].astype(str).values)) if len(info["excluded"]) else set()
    bars = []
    for _, r in props.iterrows():
        label = str(r.get("Organoid", r.get("Image_name", "")))
        if "Clone" in r and props["Clone"].nunique() > 1:
            label = f"{r['Clone']} · {label}"
        bars.append({
            "label": label,
            "group": str(r.get(group_col, "")) if group_col else "",
            "slices": int(r["Slices"]),
            "excluded": tuple(str(r[k]) for k in info["key"]) in excluded_keys,
            "living": {c["id"]: _num(r.get(f"{c['name']} (% living)")) for c in cats if c["id"] != "dead"},
            "all": {c["id"]: _num(r.get(f"{c['name']} (% all)")) for c in cats},
        })
    return {
        "excel": xlsx,
        "excel_name": os.path.basename(xlsx),
        "plots": plots,
        "n_images": int(len(per_image)),
        "n_nuclei": total,
        "n_living": living,
        "unit": info["unit"],
        "n_units": int(len(tables["per_organoid"])),
        "n_units_kept": int(len(tables["curated"])),
        "min_slices": int(s["min_slices"]),
        "group_col": group_col,
        "groups": sorted({b["group"] for b in bars if b["group"]}),
        "bars": bars,
        "per_image": json.loads(per_image.to_json(orient="records")),
        "stats": json.loads(stats_df.to_json(orient="records")) if len(stats_df) else [],
        "stat_warnings": stat_warnings + info["warnings"],
    }


def _num(v):
    try:
        v = float(v)
        return None if not np.isfinite(v) else round(v, 3)
    except (TypeError, ValueError):
        return None


def _plots(project, tables, group_col, out_dir):
    plot_dir = os.path.join(out_dir, "plots")
    os.makedirs(plot_dir, exist_ok=True)
    cats = [c for c in project.categories if c["id"] != "dead"]
    props = tables["proportions"]
    files = []
    if len(props):
        label_col = "Organoid" if "Organoid" in props else "Image_name"
        order = props.sort_values([group_col, label_col] if group_col else [label_col])
        fig, ax = plt.subplots(figsize=(max(5, 0.5 * len(order) + 2.5), 4.2))
        bottom = np.zeros(len(order))
        x = np.arange(len(order))
        for c in cats:
            v = order[f"{c['name']} (% living)"].fillna(0).to_numpy()
            ax.bar(x, v, bottom=bottom, color=c["color"], edgecolor="white", linewidth=0.5, label=c["name"])
            bottom += v
        labels = [f"{r[group_col]}\n{r[label_col]}" if group_col else str(r[label_col]) for _, r in order.iterrows()]
        ax.set_xticks(x, labels, fontsize=8)
        ax.set_ylabel("% of living cells")
        ax.set_ylim(0, 100)
        ax.spines[["top", "right"]].set_visible(False)
        ax.legend(frameon=False, bbox_to_anchor=(1.01, 1), loc="upper left", fontsize=8)
        fig.tight_layout()
        for ext in ("png", "svg"):
            p = os.path.join(plot_dir, f"proportions_living.{ext}")
            fig.savefig(p, dpi=200)
            files.append(p)
        plt.close(fig)

        if group_col and props[group_col].nunique() >= 2:
            groups = sorted(props[group_col].unique())
            fig, axes = plt.subplots(1, len(cats), figsize=(2.4 * len(cats) + 1, 3.6), squeeze=False)
            rng = np.random.default_rng(0)
            for ax, c in zip(axes[0], cats):
                col = f"{c['name']} (% living)"
                data = [props.loc[props[group_col] == g, col].dropna().to_numpy() for g in groups]
                ax.boxplot(data, widths=0.5, showfliers=False, medianprops={"color": "black"})
                for i, d in enumerate(data):
                    ax.scatter(i + 1 + rng.uniform(-0.12, 0.12, len(d)), d, s=16, color=c["color"], edgecolor="black", linewidth=0.4, zorder=3)
                ax.set_xticks(range(1, len(groups) + 1), groups, fontsize=8)
                ax.set_title(c["name"], fontsize=10)
                ax.spines[["top", "right"]].set_visible(False)
            axes[0][0].set_ylabel("% of living cells")
            fig.tight_layout()
            for ext in ("png", "svg"):
                p = os.path.join(plot_dir, f"groups_living.{ext}")
                fig.savefig(p, dpi=200)
                files.append(p)
            plt.close(fig)
    return files
